# -*- coding: utf-8 -*-
"""生成软著【软件开发基本信息表】Word 文档（对应中国版权保护中心登记申请表填报字段）。
依赖：python-docx
运行：python gen_basic_info_docx.py
所有内容都来自 ruanzhu_config.C，改配置即可换项目。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ruanzhu_config import C

OUT = os.path.join(C["output_dir"],
                   "01-软件开发基本信息表-" + C["sw_name"] + C["version"] + ".docx")
FONT = "宋体"


def set_cjk(run, size=None, bold=None):
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_kv_table(doc, rows, c0=4.5, c1=11.0):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        r = t.add_row().cells
        r[0].text = ""
        r[1].text = ""
        set_cjk(r[0].paragraphs[0].add_run(k), size=10.5, bold=True)
        set_cjk(r[1].paragraphs[0].add_run(v), size=10.5)
    for row in t.rows:
        row.cells[0].width = Cm(c0)
        row.cells[1].width = Cm(c1)
    doc.add_paragraph()
    return t


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    set_cjk(p.add_run(text), size=14 if level == 1 else 12, bold=True)
    return p


doc = Document()
# 默认正文字体
style = doc.styles["Normal"]
style.font.name = FONT
_rpr = style.element.get_or_add_rPr()
_rfonts = _rpr.find(qn('w:rFonts'))
if _rfonts is None:
    _rfonts = OxmlElement('w:rFonts')
    _rpr.append(_rfonts)
_rfonts.set(qn('w:eastAsia'), FONT)
style.font.size = Pt(10.5)

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cjk(title.add_run(C["sw_name"] + C["version"] + " 软件著作权登记申请信息表"), size=16, bold=True)
note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cjk(note.add_run("用途：本表汇总“计算机软件著作权登记申请”所需填报的基础信息，"
                     "供在中国版权保护中心（www.ccopyright.com.cn）办理登记时对照填写。"),
        size=9, bold=False)
doc.add_paragraph()

# 一、软件基本信息
add_heading(doc, "一、软件基本信息", 1)
add_kv_table(doc, [
    ("软件全称", C["sw_name"]),
    ("软件简称", C["sw_short"] or "—"),
    ("版本号", C["version"]),
    ("软件作品说明", "原创"),
    ("开发完成日期", C["dev_finish_date"]),
    ("首次发表日期", C["first_publish"]),
    ("开发方式", C["dev_method"]),
    ("权利取得方式", C["rights_method"]),
    ("权利范围", C["rights_scope"]),
    ("软件用途", C["sw_purpose"]),
])

# 二、著作权人信息
add_heading(doc, "二、著作权人信息", 1)
add_kv_table(doc, [
    ("著作权人名称", C["owner_name"]),
    ("主体类型", "企业法人"),
    ("统一社会信用代码", C["owner_credit"]),
    ("注册地址", C["owner_addr"]),
    ("国标行业分类", C["owner_industry"]),
])

# 三、软件开发环境
add_heading(doc, "三、软件开发环境", 1)
add_kv_table(doc, list(C["dev_env"].items()))

# 四、软件运行环境
add_heading(doc, "四、软件运行环境", 1)
add_kv_table(doc, list(C["run_env"].items()))

# 五、软件功能与技术特点
add_heading(doc, "五、软件功能与技术特点", 1)
add_heading(doc, "5.1 主要功能", 2)
ft = doc.add_table(rows=1, cols=3)
ft.style = "Table Grid"
hdr = ft.rows[0].cells
for i, h in enumerate(["序号", "功能模块", "功能说明"]):
    set_cjk(hdr[i].paragraphs[0].add_run(h), size=10.5, bold=True)
for no, name, desc in C["functions"]:
    r = ft.add_row().cells
    set_cjk(r[0].paragraphs[0].add_run(no), size=10)
    set_cjk(r[1].paragraphs[0].add_run(name), size=10)
    set_cjk(r[2].paragraphs[0].add_run(desc), size=10)
for row in ft.rows:
    row.cells[0].width = Cm(1.5)
    row.cells[1].width = Cm(4.5)
    row.cells[2].width = Cm(9.5)
doc.add_paragraph()

add_heading(doc, "5.2 技术特点", 2)
for feat in C["tech_features"]:
    p = doc.add_paragraph(style="List Bullet")
    set_cjk(p.add_run(feat), size=10.5)

# 六、提交材料清单
add_heading(doc, "六、提交材料清单", 1)
mt = doc.add_table(rows=1, cols=4)
mt.style = "Table Grid"
for i, h in enumerate(["序号", "材料名称", "格式", "状态"]):
    set_cjk(mt.rows[0].cells[i].paragraphs[0].add_run(h), size=10.5, bold=True)
items = [
    ("1", "软件著作权登记申请表（按本信息表填写）", "在线填报", "待填报"),
    ("2", "程序鉴别材料（源程序前 30 页 + 后 30 页，每页不少于 50 行）", "PDF", "已生成"),
    ("3", "文档鉴别材料（软件使用说明书）", "PDF", "已生成"),
    ("4", "著作权人身份证明（企业营业执照副本）", "扫描件", "待提供"),
]
for no, name, fmt, st in items:
    r = mt.add_row().cells
    set_cjk(r[0].paragraphs[0].add_run(no), size=10)
    set_cjk(r[1].paragraphs[0].add_run(name), size=10)
    set_cjk(r[2].paragraphs[0].add_run(fmt), size=10)
    set_cjk(r[3].paragraphs[0].add_run(st), size=10)
for row in mt.rows:
    row.cells[0].width = Cm(1.5)
    row.cells[1].width = Cm(9.5)
    row.cells[2].width = Cm(2.5)
    row.cells[3].width = Cm(2.5)

doc.save(OUT)
print("已生成:", OUT)
